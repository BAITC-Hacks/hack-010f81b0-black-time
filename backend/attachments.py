"""Bounded local text extraction. Uploaded text is always an unconfirmed draft."""

import io
import os
import subprocess
import tempfile
import warnings
import zipfile
from pathlib import Path

from fastapi import HTTPException

MAX_TEXT = 10000


def windows_ocr(content: bytes) -> str:
    if os.name != "nt":
        raise HTTPException(503, "JPEG мәтінін тану үшін Windows OCR қажет. Тауар атауын мәтінмен енгізіңіз.")
    from PIL import Image

    with warnings.catch_warnings():
        warnings.simplefilter("error", Image.DecompressionBombWarning)
        with Image.open(io.BytesIO(content)) as image:
            if image.format != "JPEG" or image.width * image.height > 16000000:
                raise ValueError("Unsupported or oversized image")
            image.verify()
    temp_path = None
    try:
        with tempfile.NamedTemporaryFile(suffix=".jpg", delete=False) as file:
            file.write(content)
            temp_path = Path(file.name)
        completed = subprocess.run(
            ["powershell.exe", "-NoProfile", "-NonInteractive", "-ExecutionPolicy", "Bypass", "-File", str(Path(__file__).with_name("ocr.ps1")), "-ImagePath", str(temp_path)],
            capture_output=True, text=True, encoding="utf-8", timeout=25,
            creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0),
        )
        if completed.returncode:
            raise HTTPException(503, "Windows OCR іске қосылмады. Windows-та орысша не ағылшынша OCR тілін орнатыңыз.")
        return completed.stdout.strip()
    except subprocess.TimeoutExpired:
        raise HTTPException(504, "Сурет мәтінін тану уақыты бітті.") from None
    finally:
        if temp_path:
            temp_path.unlink(missing_ok=True)


def extract_attachment(filename: str, content: bytes) -> dict:
    filename = filename.replace("\\", "/").rsplit("/", 1)[-1][:200]
    suffix = Path(filename).suffix.lower()
    if suffix not in {".xlsx", ".xls", ".docx", ".pdf", ".jpg", ".jpeg"}:
        raise HTTPException(415, "XLSX/XLS, DOCX, PDF немесе JPEG файлын таңдаңыз. Ескі .doc файлын .docx ретінде сақтаңыз.")
    if not content:
        raise HTTPException(422, "Файл бос.")
    if suffix in {".xlsx", ".docx"}:
        try:
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                if len(archive.infolist()) > 2000 or sum(item.file_size for item in archive.infolist()) > 30 * 1024 * 1024:
                    raise HTTPException(413, "Файл ашылғанда тым үлкен. Шағын құжат жіберіңіз.")
        except zipfile.BadZipFile:
            raise HTTPException(422, "Office файлының пішімі жарамсыз.") from None
    notes = ["Бұл — файлдан алынған жоба мәтіні. Қателерін тексеріп, содан кейін чатқа жіберіңіз. Файл себетке қосуға растау болып саналмайды."]
    try:
        lines = []
        if suffix == ".xlsx":
            from openpyxl import load_workbook
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True, keep_links=False)
            try:
                for sheet in workbook.worksheets[:5]:
                    for row in sheet.iter_rows(max_row=200, max_col=20, values_only=True):
                        lines.append(" | ".join(str(value) for value in row if value is not None))
            finally:
                workbook.close()
            notes.append("Алғашқы 5 парақ, әр парақтың 200 жолы және 20 бағанына дейін оқылады. Формулалар орындалмайды.")
        elif suffix == ".xls":
            import xlrd
            workbook = xlrd.open_workbook(file_contents=content, on_demand=True)
            try:
                for sheet in workbook.sheets()[:5]:
                    for index in range(min(sheet.nrows, 200)):
                        lines.append(" | ".join(str(value) for value in sheet.row_values(index)[:20] if value != ""))
            finally:
                workbook.release_resources()
        elif suffix == ".docx":
            from docx import Document
            document = Document(io.BytesIO(content))
            lines = [paragraph.text for paragraph in document.paragraphs[:500]]
            for table in document.tables[:20]:
                for row in table.rows[:200]:
                    lines.append(" | ".join(cell.text for cell in row.cells[:20]))
        elif suffix == ".pdf":
            from pypdf import PdfReader
            reader = PdfReader(io.BytesIO(content))
            if reader.is_encrypted:
                raise HTTPException(422, "Құпиясөзбен жабылған PDF оқылмайды.")
            lines = [page.extract_text() or "" for page in reader.pages[:10]]
            notes.append("PDF-тің алғашқы 10 бетінің мәтіні оқылды. Мәтін қабаты жоқ скан үшін JPEG OCR қолданыңыз.")
        else:
            lines = [windows_ocr(content)]
            notes.append("JPEG-тен тек көрінетін жазу танылады; суреттегі тауар моделі автоматты түрде расталмайды.")
        text = "\n".join(line for line in lines if line.strip()).strip()
        if not text:
            raise HTTPException(422, "Файлдан мәтін табылмады. Тауар атауын не артикулын қолмен жазыңыз.")
        if len(text) > MAX_TEXT:
            notes.append("Мәтін алғашқы 10 000 таңбаға дейін қысқартылды.")
        return {"filename": filename, "text": text[:MAX_TEXT], "notes": notes}
    except HTTPException:
        raise
    except Exception:
        # Third-party parsers can include document text in exceptions; never expose it.
        raise HTTPException(422, "Файлды оқу мүмкін болмады. Пішімін тексеріп, қайта жіберіңіз.") from None
