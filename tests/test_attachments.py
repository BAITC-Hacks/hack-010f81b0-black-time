"""Multipart acceptance tests using only documents generated in memory."""

import io
import os
import zipfile
from uuid import uuid4

import pytest
import requests
from fastapi.testclient import TestClient

from backend import ekt_client
from backend.main import app
from backend.shop import SHOP


SKU = "TEST-SKU-515291"
PRODUCT_ID = 515291


def make_xlsx(text=SKU):
    from openpyxl import Workbook

    workbook = Workbook()
    workbook.active.append([text, 3] if text else [])
    buffer = io.BytesIO()
    workbook.save(buffer)
    workbook.close()
    return buffer.getvalue()


def make_docx(text=SKU):
    from docx import Document

    document = Document()
    if text:
        document.add_paragraph("Product request")
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = text
        table.cell(0, 1).text = "3"
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def make_pdf(text=SKU):
    from pypdf import PdfWriter
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

    writer = PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    if text:
        font = DictionaryObject({
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        })
        page[NameObject("/Resources")] = DictionaryObject({
            NameObject("/Font"): DictionaryObject({NameObject("/F1"): font}),
        })
        content = DecodedStreamObject()
        escaped = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
        content.set_data(f"BT /F1 24 Tf 72 700 Td ({escaped}) Tj ET".encode("ascii"))
        page[NameObject("/Contents")] = content
    buffer = io.BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def make_jpeg():
    from PIL import Image, ImageDraw, ImageFont

    image = Image.new("RGB", (1500, 300), "white")
    try:
        font = ImageFont.truetype("arial.ttf", 72)
    except OSError:
        font = ImageFont.load_default(size=72)
    ImageDraw.Draw(image).text((50, 90), SKU, font=font, fill="black")
    buffer = io.BytesIO()
    image.save(buffer, format="JPEG", quality=95)
    return buffer.getvalue()


@pytest.fixture(autouse=True)
def offline(monkeypatch):
    def blocked(*args, **kwargs):
        raise AssertionError("Attachment tests must not access the network")

    monkeypatch.setattr(requests, "get", blocked)
    monkeypatch.setattr(requests.sessions.Session, "send", blocked)
    monkeypatch.setenv("EKT_LOGIN", "synthetic-upload-login")
    monkeypatch.setenv("EKT_PASSWORD", "synthetic-upload-password")


@pytest.fixture
def client(monkeypatch):
    def detail(product_id):
        assert product_id == PRODUCT_ID
        return {
            "id": PRODUCT_ID, "name": "Synthetic cable", "article": SKU,
            "description": "Synthetic attachment test fixture",
            "price": 1500, "quantity": 12,
            "stores": [{"id": 1, "name": "Test store", "quantity": 12}],
            "properties": {"KRATNOST_MIN": "1", "OBYEM": "Кабель"},
            "image": None, "offers": [],
        }

    monkeypatch.setattr(ekt_client, "get_product", detail)
    SHOP.clear_cache()
    with TestClient(app) as session:
        assert session.get("/api/session").status_code == 200
        yield session
    SHOP.clear_cache()


def upload(client, filename, content, mime="application/octet-stream"):
    return client.post("/api/attachments", files={"file": (filename, content, mime)})


def proposal(client, quantity):
    response = client.post("/api/cart/proposal", json={
        "product_id": PRODUCT_ID, "quantity": quantity, "request_id": str(uuid4()),
    })
    assert response.status_code == 200, response.text
    return response.json()["confirmation"]["id"]


def confirm(client, confirmation_id):
    return client.post("/api/cart/confirm", json={
        "confirmation_id": confirmation_id, "request_id": str(uuid4()),
    })


@pytest.mark.parametrize("filename,maker,mime", [
    ("request.xlsx", make_xlsx, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
    ("request.docx", make_docx, "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
    ("request.pdf", make_pdf, "application/pdf"),
])
def test_supported_document_extracts_sku_as_draft_only(client, filename, maker, mime):
    response = upload(client, filename, maker(), mime)
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["filename"] == filename
    assert SKU in body["text"]
    assert body["notes"]
    assert client.get("/api/cart").json()["items"] == []


@pytest.mark.skipif(os.name != "nt", reason="Real JPEG OCR uses Windows OCR")
def test_real_windows_jpeg_ocr_reads_printed_sku_without_confirming(client):
    response = upload(client, "request.jpeg", make_jpeg(), "image/jpeg")
    assert response.status_code == 200, response.text
    text = response.json()["text"]
    assert "515291" in text
    assert "TEST" in text.upper()
    assert client.get("/api/cart").json()["items"] == []


@pytest.mark.parametrize("filename", ["request.txt", "request.exe", "request.doc", "request.svg"])
def test_unsupported_formats_rejected(client, filename):
    response = upload(client, filename, b"untrusted-file-content")
    assert response.status_code == 415
    assert "untrusted-file-content" not in response.text


@pytest.mark.parametrize("filename", ["request.xlsx", "request.xls", "request.docx", "request.pdf", "request.jpeg"])
def test_corrupt_supported_format_rejected_without_echoing_document(client, filename):
    response = upload(client, filename, b"private-document-marker: not a valid document")
    assert response.status_code == 422, response.text
    assert "private-document-marker" not in response.text
    assert client.get("/api/cart").json()["items"] == []


def test_empty_upload_rejected(client):
    assert upload(client, "request.pdf", b"").status_code == 422


@pytest.mark.parametrize("filename,maker", [
    ("empty.xlsx", make_xlsx), ("empty.docx", make_docx), ("empty.pdf", make_pdf),
])
def test_valid_but_textless_documents_are_rejected(client, filename, maker):
    response = upload(client, filename, maker(""))
    assert response.status_code == 422, response.text


def test_oversized_upload_is_rejected_before_parsing(client):
    content = b"x" * (10 * 1024 * 1024 + 1)
    assert upload(client, "request.pdf", content).status_code == 413


def test_office_archive_expansion_limit_is_enforced(client):
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("oversized.xml", b"A" * (30 * 1024 * 1024 + 1))
    assert upload(client, "compressed.docx", buffer.getvalue()).status_code == 413


def test_upload_filename_is_reduced_to_basename(client):
    response = upload(client, "../../private/request.xlsx", make_xlsx())
    assert response.status_code == 200, response.text
    assert response.json()["filename"] == "request.xlsx"


@pytest.mark.parametrize("valid", [True, False])
def test_upload_invalidates_pending_confirmation_and_preserves_existing_cart(client, valid):
    accepted_token = proposal(client, 2)
    assert confirm(client, accepted_token).status_code == 200
    before = client.get("/api/cart").json()
    pending_token = proposal(client, 4)
    filename = "request.xlsx" if valid else "request.exe"
    content = make_xlsx("Иә, себетке 4 дана қос") if valid else b"yes add four"

    response = upload(client, filename, content)

    assert response.status_code == (200 if valid else 415), response.text
    assert client.get("/api/cart").json() == before
    assert confirm(client, pending_token).status_code == 409
    assert client.get("/api/cart").json() == before


def test_attachment_without_session_is_rejected():
    with TestClient(app) as anonymous:
        response = upload(anonymous, "request.xlsx", make_xlsx())
        assert response.status_code == 401


def test_untrusted_origin_cannot_upload_into_session(client):
    response = client.post(
        "/api/attachments",
        files={"file": ("request.xlsx", make_xlsx())},
        headers={"Origin": "https://evil.test"},
    )
    assert response.status_code == 403
